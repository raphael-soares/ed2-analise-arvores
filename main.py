import argparse
import os
import statistics
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

RAIZ = Path(__file__).resolve().parent
SRC = RAIZ / "src"
BUILD = RAIZ / "build"
EXECUTAVEL = BUILD / "arvores"

ESTRUTURAS = ["avl", "rn", "b1", "b5", "b10"]
ROTULOS = {"avl": "AVL", "rn": "Rubro-negra", "b1": "B ordem 1", "b5": "B ordem 5", "b10": "B ordem 10"}
COLUNAS = [
    "n",
    "adicao_busca",
    "adicao_ajuste",
    "adicao_total",
    "adicao_desvio",
    "remocao_busca",
    "remocao_ajuste",
    "remocao_total",
    "remocao_desvio",
]
CORES = {
    "AVL": "tab:blue",
    "Rubro-negra": "tab:red",
    "B ordem 1": "tab:green",
    "B ordem 5": "tab:orange",
    "B ordem 10": "tab:purple",
}
INDIVIDUAIS = [
    ("avl", "AVL", ["AVL"]),
    ("rn", "Rubro-negra", ["Rubro-negra"]),
    ("b", "árvore B", ["B ordem 1", "B ordem 5", "B ordem 10"]),
]


def compilar():
    BUILD.mkdir(exist_ok=True)
    fontes = sorted(str(p) for p in SRC.glob("*.c"))
    comando = ["gcc", "-std=c11", "-O2", "-Wall", "-Wextra", "-o", str(EXECUTAVEL), *fontes]
    if subprocess.run(comando).returncode != 0:
        sys.exit(1)


def garantir_compilado():
    if not EXECUTAVEL.exists():
        compilar()
        return
    compilado = EXECUTAVEL.stat().st_mtime
    if any(p.stat().st_mtime > compilado for p in SRC.iterdir()):
        compilar()


def executar(estrutura, semente, tamanhos):
    comando = [str(EXECUTAVEL), estrutura, str(semente), *map(str, tamanhos)]
    resultado = subprocess.run(comando, capture_output=True, text=True)
    if resultado.returncode != 0:
        sys.stderr.write(resultado.stderr)
        sys.exit(1)
    return [tuple(map(int, linha.split(","))) for linha in resultado.stdout.split()]


def media(valores):
    return round(statistics.fmean(valores), 2)


def desvio(valores):
    return round(statistics.stdev(valores), 2) if len(valores) > 1 else 0


def resumir(execucoes):
    linhas = []
    for amostras in zip(*execucoes):
        ad_busca = [a[1] for a in amostras]
        ad_ajuste = [a[2] for a in amostras]
        rm_busca = [a[3] for a in amostras]
        rm_ajuste = [a[4] for a in amostras]
        ad_total = [b + a for b, a in zip(ad_busca, ad_ajuste)]
        rm_total = [b + a for b, a in zip(rm_busca, rm_ajuste)]
        linhas.append([
            amostras[0][0],
            media(ad_busca),
            media(ad_ajuste),
            media(ad_total),
            desvio(ad_total),
            media(rm_busca),
            media(rm_ajuste),
            media(rm_total),
            desvio(rm_total),
        ])
    return linhas


def positivo(texto):
    valor = int(texto)
    if valor < 1:
        raise argparse.ArgumentTypeError("precisa ser maior que zero")
    return valor


def cmd_compilar(args):
    compilar()


def cmd_rodar(args):
    garantir_compilado()
    _, ad_busca, ad_ajuste, rm_busca, rm_ajuste = executar(args.estrutura, args.semente, [args.n])[0]
    print(f"{args.estrutura}  n={args.n}  semente={args.semente}")
    print(f"adição   busca={ad_busca}  ajuste={ad_ajuste}  total={ad_busca + ad_ajuste}")
    print(f"remoção  busca={rm_busca}  ajuste={rm_ajuste}  total={rm_busca + rm_ajuste}")


def cmd_experimento(args):
    import openpyxl
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font

    inicio = time.perf_counter()
    garantir_compilado()
    estruturas = [e for e in ESTRUTURAS if e in args.estruturas]
    grade = sorted({1} | set(range(args.passo, args.n_max + 1, args.passo)))
    sementes = range(1, args.amostras + 1)

    with ThreadPoolExecutor(max_workers=os.cpu_count()) as pool:
        jobs = {(e, s): pool.submit(executar, e, s, grade) for e in estruturas for s in sementes}
        resumos = {e: resumir([jobs[e, s].result() for s in sementes]) for e in estruturas}

    livro = openpyxl.Workbook()
    livro.remove(livro.active)
    for titulo, coluna in (("Adição", COLUNAS.index("adicao_total")), ("Remoção", COLUNAS.index("remocao_total"))):
        aba = livro.create_sheet(titulo)
        aba.append(["n", *(ROTULOS[e] for e in estruturas)])
        for i, n in enumerate(grade):
            aba.append([n, *(resumos[e][i][coluna] for e in estruturas)])

        grafico = LineChart()
        grafico.title = f"{titulo}: esforço médio"
        grafico.x_axis.title = "Tamanho do conjunto (n)"
        grafico.y_axis.title = "Esforço (operações)"
        grafico.x_axis.delete = False
        grafico.y_axis.delete = False
        grafico.width = 24
        grafico.height = 12
        ultima = len(grade) + 1
        grafico.add_data(Reference(aba, min_col=2, max_col=len(estruturas) + 1, min_row=1, max_row=ultima), titles_from_data=True)
        grafico.set_categories(Reference(aba, min_col=1, min_row=2, max_row=ultima))
        for serie in grafico.series:
            serie.marker.symbol = "none"
            serie.smooth = False
        aba.add_chart(grafico, "H2")

    for estrutura in estruturas:
        aba = livro.create_sheet(ROTULOS[estrutura])
        aba.append(COLUNAS)
        for linha in resumos[estrutura]:
            aba.append(linha)

    aba = livro.create_sheet("Sobre")
    for linha in [
        ("campo", "valor"),
        ("busca", "comparações e passos de descida"),
        ("ajuste", "rotações, recolorações, nós visitados no balanceamento e chaves movidas"),
        ("total", "busca + ajuste"),
        ("amostras", args.amostras),
        ("n_max", args.n_max),
        ("passo", args.passo),
        ("sementes", f"1 a {args.amostras}"),
        ("comando", " ".join(sys.argv)),
    ]:
        aba.append(linha)

    for aba in livro.worksheets:
        for celula in aba[1]:
            celula.font = Font(bold=True)
        aba.freeze_panes = "A2"

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    livro.save(saida / "esforco.xlsx")
    print(f"planilha: {saida / 'esforco.xlsx'}")
    print(f"tempo: {time.perf_counter() - inicio:.1f} s")


def ler_aba(livro, nome):
    linhas = list(livro[nome].iter_rows(values_only=True))
    ns = [linha[0] for linha in linhas[1:]]
    series = {rotulo: [linha[i] for linha in linhas[1:]] for i, rotulo in enumerate(linhas[0]) if i > 0}
    return ns, series


def desenhar(plt, ns, series, rotulos, titulo, destino, eixo_y="Esforço (operações)"):
    fig, ax = plt.subplots(figsize=(8, 5))
    for rotulo in rotulos:
        ax.plot(ns, series[rotulo], label=rotulo, color=CORES[rotulo])
    ax.set_title(titulo)
    ax.set_xlabel("Tamanho do conjunto (n)")
    ax.set_ylabel(eixo_y)
    ax.grid(True)
    ax.legend()
    fig.tight_layout()
    fig.savefig(destino.with_suffix(".png"), dpi=200)
    fig.savefig(destino.with_suffix(".pdf"))
    plt.close(fig)
    print(destino.with_suffix(".png"))


def cmd_graficos(args):
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import openpyxl

    entrada = Path(args.entrada)
    if not entrada.exists():
        print(f"arquivo não encontrado: {entrada}", file=sys.stderr)
        sys.exit(1)
    livro = openpyxl.load_workbook(entrada, read_only=True)
    operacoes = [("adicao", "Adição", *ler_aba(livro, "Adição")), ("remocao", "Remoção", *ler_aba(livro, "Remoção"))]
    livro.close()

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    todas = [ROTULOS[e] for e in ESTRUTURAS]
    completo = all(rotulo in operacoes[0][3] for rotulo in todas)
    if not completo:
        print("comparativo exige as 5 estruturas")

    for prefixo, operacao, ns, series in operacoes:
        if completo:
            desenhar(plt, ns, series, todas, f"{operacao}: comparativo", saida / prefixo)
            por_operacao = {r: [v / n for v, n in zip(series[r], ns)] for r in todas}
            desenhar(
                plt,
                ns,
                por_operacao,
                todas,
                f"{operacao}: esforço por operação",
                saida / f"{prefixo}-por-operacao",
                "Esforço médio por operação",
            )
        for nome, titulo, rotulos in INDIVIDUAIS:
            if all(rotulo in series for rotulo in rotulos):
                desenhar(plt, ns, series, rotulos, f"{operacao}: {titulo}", saida / f"{prefixo}-{nome}")


def main():
    parser = argparse.ArgumentParser(description="Esforço de adição e remoção em árvores balanceadas.")
    sub = parser.add_subparsers(dest="comando", required=True)

    p = sub.add_parser("compilar", help="compila o executável em build/")
    p.set_defaults(func=cmd_compilar)

    p = sub.add_parser("rodar", help="roda uma estrutura com um conjunto")
    p.add_argument("estrutura", choices=ESTRUTURAS)
    p.add_argument("--n", type=int, default=1000)
    p.add_argument("--semente", type=int, default=1)
    p.set_defaults(func=cmd_rodar)

    p = sub.add_parser("experimento", help="mede todas as estruturas e grava a planilha")
    p.add_argument("--estruturas", nargs="+", choices=ESTRUTURAS, default=ESTRUTURAS)
    p.add_argument("--n-max", type=positivo, default=10000)
    p.add_argument("--passo", type=positivo, default=1)
    p.add_argument("--amostras", type=positivo, default=10)
    p.add_argument("--saida", default="resultados")
    p.set_defaults(func=cmd_experimento)

    p = sub.add_parser("graficos", help="gera os gráficos a partir da planilha")
    p.add_argument("--entrada", default="resultados/esforco.xlsx")
    p.add_argument("--saida", default="graficos")
    p.set_defaults(func=cmd_graficos)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
